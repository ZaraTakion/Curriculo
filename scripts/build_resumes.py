#!/usr/bin/env python3
"""Build bilingual DOCX and PDF resumes from Markdown sources (no paid services)."""
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
BLUE = RGBColor(19, 58, 102)
ACCENT = RGBColor(14, 113, 139)
INK = RGBColor(32, 48, 63)
GREY = RGBColor(83, 98, 109)
TOKEN = re.compile(r"(\*\*|\[[^\]]+\]\(https?://[^)]+\)|\[[^\]]+\]\(mailto:[^)]+\))")


def hyperlink(paragraph, caption: str, target: str, bold: bool) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0E718B")
    props.append(color)
    if bold:
        props.append(OxmlElement("w:b"))
    run.append(props)
    value = OxmlElement("w:t")
    value.text = caption
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, value: str) -> None:
    strong = False
    for part in TOKEN.split(value):
        if not part:
            continue
        if part == "**":
            strong = not strong
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if match:
            hyperlink(paragraph, match.group(1), match.group(2), strong)
        else:
            run = paragraph.add_run(part)
            run.bold = strong


def configure(document) -> None:
    sec = document.sections[0]
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.60)
    sec.left_margin = Inches(0.73)
    sec.right_margin = Inches(0.73)
    sec.header_distance = Inches(0.24)
    sec.footer_distance = Inches(0.30)

    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10.1)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.line_spacing = 1.18
    pf.space_after = Pt(4)

    for name, size, before, after, color in (
        ("Title", 19.5, 0, 4, BLUE),
        ("Heading 2", 11.5, 10, 3, BLUE),
    ):
        style = document.styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def generate(source: Path, target: Path) -> None:
    lines = [s.strip() for s in source.read_text(encoding="utf-8").splitlines()]
    doc = Document()
    configure(doc)
    doc.core_properties.author = "Rodrigo Araújo Maciel Pinheiro"
    doc.core_properties.title = "Currículo — Back-End Python" if "pt" in source.stem else "Resume — Python Back-End"

    for line in lines:
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(p, line[2:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            inline(p, line[3:])
        elif line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.10)
            p.paragraph_format.first_line_indent = Inches(-0.10)
            p.add_run("• ").font.color.rgb = ACCENT
            inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            if line.startswith("**") and line.endswith("**"):
                p.paragraph_format.space_after = Pt(5)
            inline(p, line)
            if "mailto:" in line:
                p.paragraph_format.space_after = Pt(7)
                for r in p.runs:
                    r.font.size = Pt(8.3)
                    r.font.color.rgb = GREY
    doc.save(target)


def main() -> None:
    if shutil.which("libreoffice") is None:
        raise RuntimeError("LibreOffice Writer required to export PDFs")
    for language, source in (("PTBR", "curriculo_pt.md"), ("EN", "resume_en.md")):
        docx_path = ROOT / f"Rodrigo_Pinheiro_2026_{language}.docx"
        pdf_path = docx_path.with_suffix(".pdf")
        generate(ROOT / "src" / source, docx_path)
        command = [
            "libreoffice", "-env:UserInstallation=file:///tmp/lo-curriculo-build",
            "--headless", "--convert-to", "pdf:writer_pdf_Export",
            "--outdir", str(ROOT), str(docx_path)
        ]
        outcome = subprocess.run(command, capture_output=True, text=True, check=False)
        if outcome.returncode or not pdf_path.is_file():
            raise RuntimeError(f"Failed to export {docx_path.name}: {outcome.stdout} {outcome.stderr}")
        print(f"[BUILD] {docx_path.name} / {pdf_path.name}")


if __name__ == "__main__":
    main()
